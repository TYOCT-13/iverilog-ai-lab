# -*- coding: utf-8 -*-
"""ICARUS 答辩 PPT 重构 —— 内层布局引擎（配色/网格/组件）"""
import copy
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.oxml.ns import qn
from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# ---------- 设计令牌 ----------
BG = "F5F7F8"
BG_DARK = "1C2023"
INK = "182025"          # 主文字
GRAY = "53636B"         # 次级文字
GRAY2 = "75858D"        # 图表常态色 / 弱化
ACCENT = "00ABC8"       # 品牌青
ACCENT_SOFT = "E4F4F8"  # 高亮卡片底
RULE = "C9D3D8"         # 分隔线 / 表线
WHITE = "FFFFFF"

FONT = "Noto Sans SC"
FONT_NUM = "Bahnschrift"

SW, SH = 12192000, 6858000
L = 609600
CW = 10972800
R = L + CW                       # 11582400
TITLE_T, TITLE_H = 419100, 638175
PAGENO_X, PAGENO_Y, PAGENO_W, PAGENO_H = 11077575, 447675, 571500, 495300
RULE_Y = 1209675
FOOT_T = 6076950
PANEL_T, PANEL_H = 1450000, 3930000   # bottom 5380000


def rgb(h):
    return RGBColor.from_string(h)


def _nofit(tf):
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0


def _run(p, text, size, bold=False, color=INK, font=FONT, spc=None, italic=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    # 中英文统一字体
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = etree.SubElement(rPr, qn(tag))
        e.set("typeface", font)
    if spc is not None:
        rPr.set("spc", str(int(spc * 100)))
    return r


def text(slide, l, t, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    """paras: 每项 {runs:[(txt, style_dict)], align, space_before, space_after, line}"""
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    _nofit(tf)
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        if spec.get("line"):
            p.line_spacing = spec["line"]
        if spec.get("space_before") is not None:
            p.space_before = Pt(spec["space_before"])
        if spec.get("space_after") is not None:
            p.space_after = Pt(spec["space_after"])
        for txt, st in spec["runs"]:
            _run(p, txt, **st)
    return box


def line(slide, x1, y1, x2, y2, color=RULE, w=0.9):
    from pptx.enum.shapes import MSO_CONNECTOR
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(y1), Emu(x2), Emu(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(w)
    return ln


def rect(slide, l, t, w, h, fill=WHITE, border=None, bw=0.9, radius=None):
    shape_t = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_t, Emu(l), Emu(t), Emu(w), Emu(h))
    if radius:
        sh.adjustments[0] = radius
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if border:
        sh.line.color.rgb = rgb(border)
        sh.line.width = Pt(bw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    _nofit(tf)
    return sh


# ---------- 页面骨架 ----------
def new_slide(prs, blank, bg=BG):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    return s


def est_w(txt, size):
    """粗略估算文本宽度（EMU），用于标题自动缩号与换行预判。"""
    units = 0.0
    for ch in txt:
        o = ord(ch)
        if 0x4E00 <= o <= 0x9FFF or 0x3000 <= o <= 0x303F or 0xFF00 <= o <= 0xFF65:
            units += 1.0
        elif ch == " ":
            units += 0.30
        elif ch.isdigit() or ("A" <= ch <= "Z"):
            units += 0.58
        elif ch.islower():
            units += 0.52
        else:
            units += 0.40
    return int(units * size * 12700)


def head(slide, title, num, dark=False):
    tcol = WHITE if dark else INK
    avail = CW - 780000          # 让开右上角页码
    size = 34.5
    for cand in (34.5, 32, 30, 28, 26):
        size = cand
        if est_w(title, cand) <= avail:
            break
    text(slide, L, TITLE_T, avail, TITLE_H,
         [{"runs": [(title, dict(size=size, bold=True, color=tcol))]}])
    text(slide, PAGENO_X, PAGENO_Y, PAGENO_W, PAGENO_H,
         [{"runs": [(f"{num:02d}", dict(size=24, bold=True, color=ACCENT, font=FONT_NUM))],
           "align": PP_ALIGN.RIGHT}])
    line(slide, L, RULE_Y, R, RULE_Y)


def foot(slide, lines, size=15.5):
    if isinstance(lines, str):
        lines = [lines]
    text(slide, L, FOOT_T, CW, 609600,
         [{"runs": [(s, dict(size=size, color=GRAY))], "space_before": 0} for s in lines])


def panel(slide, t=PANEL_T, h=PANEL_H, fill=WHITE):
    return rect(slide, L, t, CW, h, fill=fill)


def subtitle(slide, txt, t=1600000, size=22, color=INK, x=L + 190500, w=None, rule=True):
    w = w or CW - 381000
    text(slide, x, t, w, 420000,
         [{"runs": [(txt, dict(size=size, bold=True, color=color))]}])
    if rule:
        line(slide, x, t + 470000, x + w, t + 470000, RULE, 0.75)


def caption(slide, paras, t, x=L + 190500, w=None, size=15.5, color=INK, bold=False):
    w = w or CW - 381000
    if isinstance(paras, str):
        paras = [paras]
    text(slide, x, t, w, 900000,
         [{"runs": [(s, dict(size=size, bold=bold, color=color))],
           "line": 1.35, "space_after": 3} for s in paras])


def cards(slide, items, t, h=1700000, x=L + 190500, total_w=CW - 381000,
          gap=250000, title_size=20, line_size=13, num_size=15.5,
          arrow=True, fill=WHITE, border=RULE):
    n = len(items)
    w = (total_w - gap * (n - 1)) // n
    for i, it in enumerate(items):
        cx = x + i * (w + gap)
        hl = it.get("hl", False)
        rect(slide, cx, t, w, h, fill=ACCENT_SOFT if hl else fill,
             border=None if hl else border, bw=0.9)
        inner = cx + 190500
        iw = w - 381000
        ys = t + 150000
        text(slide, inner, ys, iw, 300000,
             [{"runs": [(it["no"], dict(size=num_size, bold=True, color=ACCENT, font=FONT_NUM))]}])
        text(slide, inner, ys + 340000, iw, 340000,
             [{"runs": [(it["title"], dict(size=title_size, bold=True, color=INK))]}])
        paras = []
        for s_ in it["lines"]:
            paras.append({"runs": [(s_, dict(size=line_size, color=GRAY))], "space_after": 6, "line": 1.45})
        text(slide, inner, ys + 730000, iw, h - 850000, paras)
        if arrow and i < n - 1:
            text(slide, cx + w, t + h // 2 - 180000, gap, 360000,
                 [{"runs": [("➜", dict(size=17, bold=True, color=ACCENT))], "align": PP_ALIGN.CENTER}])


def row_need(size):
    """单行文本在给定字号下的最小行高（含上下内边距），用于预估表格真实高度。"""
    return int(size * 12700 * 1.18) + 128000


def table(slide, l, t, w, rows, col_w, row_h=None, size=20.25, hdr_size=20.25,
          first_col_bold=False, aligns=None, hdr_fill=BG_DARK):
    """返回表格实际底边 y，便于后续元素定位。"""
    nr, nc = len(rows), len(rows[0])
    src = row_h or [0] * nr
    heights = [max(src[i], row_need(hdr_size if i == 0 else size)) for i in range(nr)]
    shp = slide.shapes.add_table(nr, nc, Emu(l), Emu(t), Emu(w), Emu(sum(heights)))
    tbl = shp.table
    tbl.first_row = False
    tbl.horz_banding = False
    # 列宽
    tot = sum(col_w)
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Emu(int(w * cw / tot))
    for ri, rdata in enumerate(rows):
        tbl.rows[ri].height = Emu(heights[ri])
        for ci, val in enumerate(rdata):
            cell = tbl.cell(ri, ci)
            hdr = (ri == 0)
            fill = hdr_fill if hdr else WHITE
            color = WHITE if hdr else INK
            bold = hdr or (first_col_bold and ci == 0)
            _cell(cell, val, fill, color, size if not hdr else hdr_size, bold,
                  (aligns or [PP_ALIGN.LEFT] * nc)[ci])
    return t + sum(heights)


def _cell(cell, raw, fill, color, size, bold, align=PP_ALIGN.LEFT,
          border=RULE, margins=(160020, 160020, 64000, 64000)):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        tcPr.remove(child)
    tcPr.set("marL", str(margins[0]))
    tcPr.set("marR", str(margins[1]))
    tcPr.set("marT", str(margins[2]))
    tcPr.set("marB", str(margins[3]))
    for side in ("L", "R", "T", "B"):
        ln = etree.SubElement(tcPr, qn("a:ln" + side))
        ln.set("w", "9525")
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr"))
        c.set("val", border)
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "solid")
    sf = etree.SubElement(tcPr, qn("a:solidFill"))
    c = etree.SubElement(sf, qn("a:srgbClr"))
    c.set("val", fill)
    # 文本
    raw = str(raw) if raw is not None else ""
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if raw:
        _run(p, raw, size=size, bold=bold, color=color)
    return cell


def kpi(slide, l, t, w, big, label_lines, big_size=54, label_size=17, color=ACCENT):
    text(slide, l, t, w, big_size * 14000,
         [{"runs": [(big, dict(size=big_size, bold=True, color=color, font=FONT_NUM))]}])
    paras = []
    for s_ in label_lines:
        paras.append({"runs": [(s_, dict(size=label_size, color=INK))], "space_after": 4, "line": 1.3})
    text(slide, l, t + int(big_size * 16500), w, 900000, paras)


def bar_chart(slide, l, t, w, h, cats, vals, hi=None, vmax=None, unit=None,
              dlabel_size=20, cat_size=17, ax_size=15, series_name="数值"):
    from pptx.chart.data import CategoryChartData
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(series_name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Emu(l), Emu(t), Emu(w), Emu(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.name = FONT
    ch.font.size = Pt(ax_size)
    ch.font.color.rgb = rgb(INK)
    plot = ch.plots[0]
    plot.gap_width = 62
    plot.vary_by_categories = False
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(GRAY2)
    ser.format.line.fill.background()
    if hi is not None:
        pt = ser.points[hi]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(ACCENT)
        pt.format.line.fill.background()
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.number_format = "0"
    dl.number_format_is_linked = False
    dl.font.size = Pt(dlabel_size)
    dl.font.bold = True
    dl.font.name = FONT
    dl.font.color.rgb = rgb(INK)
    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.format.line.fill.background()
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(cat_size)
    ca.tick_labels.font.name = FONT
    ca.tick_labels.font.color.rgb = rgb(INK)
    va = ch.value_axis
    va.minimum_scale = 0
    if vmax:
        va.maximum_scale = vmax
    if unit:
        va.major_unit = unit
    va.has_major_gridlines = True
    gl = va.major_gridlines.format.line
    gl.color.rgb = rgb(RULE)
    gl.width = Pt(0.6)
    va.format.line.fill.background()
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.tick_labels.font.size = Pt(ax_size)
    va.tick_labels.font.name = FONT
    va.tick_labels.font.color.rgb = rgb(INK)
    return ch


def picture(slide, path, l, t, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Emu(int(w))
    if h:
        kw["height"] = Emu(int(h))
    return slide.shapes.add_picture(path, Emu(int(l)), Emu(int(t)), **kw)


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt
